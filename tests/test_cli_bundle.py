import hashlib
import json
import re
from pathlib import Path

import pytest
from cryptography.fernet import Fernet
from docx import Document

from pii_redactor.application import (
    BundleServices,
    CreateRedactionBundle,
    RedactDocument,
)
from pii_redactor.cli import build_parser, main
from pii_redactor.docx_adapter import DocxRedactor
from pii_redactor.domain import PseudonymizedText, TextReplacement
from pii_redactor.mapping_store import EncryptedMappingStore, MappingKeyLoader
from pii_redactor.publishing import AtomicBundlePublisher
from pii_redactor.validation import DocxValidationSuite


ORIGINAL = "private@example.com"
CONFIG = {
    "config_version": "1",
    "model_version": "en_core_web_lg-3.8.0",
}


def test_cli_publishes_docx_encrypted_mapping_and_pii_safe_audit(tmp_path):
    source_path = tmp_path / "source.docx"
    config_path = tmp_path / "config.json"
    key_path = tmp_path / "mapping.key"
    output_directory = tmp_path / "bundle"
    document = Document()
    document.add_paragraph(f"Contact {ORIGINAL}")
    document.save(source_path)
    config_path.write_text(json.dumps(CONFIG), encoding="utf-8")
    key_path.write_text(Fernet.generate_key().decode(), encoding="ascii")
    source_hash = _sha256(source_path)

    exit_code = main(
        _arguments(source_path, output_directory, config_path, key_path),
        application_factory=_application_factory,
    )

    assert exit_code == 0
    output_path = output_directory / "redacted_prospectus.docx"
    mapping_path = output_directory / "redaction_mapping.json.enc"
    audit_path = output_directory / "audit_report.json"
    assert {item.name for item in output_directory.iterdir()} == {
        output_path.name,
        mapping_path.name,
        audit_path.name,
    }
    assert _sha256(source_path) == source_hash

    key = MappingKeyLoader.from_file(key_path)
    mapping = EncryptedMappingStore().load(mapping_path, key)
    audit = json.loads(audit_path.read_text(encoding="utf-8"))
    replacement = mapping.mappings[0].replacement
    assert mapping.metadata.source_sha256 == source_hash
    assert mapping.metadata.output_sha256 == _sha256(output_path)
    assert audit["source_sha256"] == source_hash
    assert audit["output_sha256"] == _sha256(output_path)
    assert audit["entity_counts"] == {"EMAIL_ADDRESS": 1}
    assert audit["status"] == "success"
    assert ORIGINAL not in audit_path.read_text(encoding="utf-8")
    assert replacement not in audit_path.read_text(encoding="utf-8")
    assert ORIGINAL not in "\n".join(
        paragraph.text for paragraph in Document(output_path).paragraphs
    )
    assert "PII_REDACTOR_MAPPING_KEY" in build_parser().format_help()
    assert "non-zero" in build_parser().format_help()


@pytest.mark.parametrize(
    "failure_point",
    (
        "missing_key",
        "invalid_key",
        "detector",
        "encryption",
        "validation",
        "publish",
    ),
)
def test_cli_failure_is_nonzero_and_leaves_no_partial_bundle(
    tmp_path, failure_point, capsys
):
    source_path = tmp_path / "source.docx"
    config_path = tmp_path / "config.json"
    key_path = tmp_path / "mapping.key"
    output_directory = tmp_path / "bundle"
    document = Document()
    document.add_paragraph(f"Contact {ORIGINAL}")
    document.save(source_path)
    config_path.write_text(json.dumps(CONFIG), encoding="utf-8")
    key_path.write_text(Fernet.generate_key().decode(), encoding="ascii")
    factory_called = False

    def failing_factory(config, seed):
        nonlocal factory_called
        factory_called = True
        return _application_factory(config, seed, failure_point)

    arguments = _arguments(source_path, output_directory, config_path, key_path)
    environment = None
    if failure_point == "missing_key":
        arguments = arguments[:-2]
        environment = {}
    elif failure_point == "invalid_key":
        key_path.write_text("not-a-fernet-key", encoding="ascii")

    exit_code = main(
        arguments,
        application_factory=failing_factory,
        environment=environment,
    )

    assert exit_code != 0
    assert not output_directory.exists()
    assert not tuple(tmp_path.glob(".bundle.*"))
    assert ORIGINAL not in capsys.readouterr().err
    if failure_point in {"missing_key", "invalid_key"}:
        assert factory_called is False


def _arguments(source, output_directory, config, key) -> list[str]:
    return [
        "--input",
        str(source),
        "--output-dir",
        str(output_directory),
        "--config",
        str(config),
        "--seed",
        "17",
        "--key-file",
        str(key),
    ]


def _application_factory(config, seed, failure_point=None):
    pseudonymizer = _EmailPseudonymizer(fail=failure_point == "detector")
    mapping_store = (
        _FailingMappingStore()
        if failure_point == "encryption"
        else EncryptedMappingStore()
    )
    validator = (
        _FailingValidator()
        if failure_point == "validation"
        else DocxValidationSuite()
    )
    move = _fail if failure_point == "publish" else None
    publisher = AtomicBundlePublisher(move=move)
    return CreateRedactionBundle(
        config=config,
        services=BundleServices(
            redact_document=RedactDocument(
                document_redactor=DocxRedactor(),
                text_pseudonymizer=pseudonymizer,
                seed=seed,
            ),
            mapping_store=mapping_store,
            validator=validator,
            publisher=publisher,
        ),
    )


class _EmailPseudonymizer:
    def __init__(self, fail=False):
        self._fail = fail

    def pseudonymize(self, text, registry):
        if self._fail:
            raise RuntimeError("detector failed")
        matches = tuple(re.finditer(re.escape(ORIGINAL), text))
        replacements = tuple(
            TextReplacement(
                match.start(),
                match.end(),
                registry.replacement_for("EMAIL_ADDRESS", match.group()),
            )
            for match in matches
        )
        transformed = text
        for replacement in reversed(replacements):
            transformed = (
                transformed[: replacement.start]
                + replacement.value
                + transformed[replacement.end :]
            )
        return PseudonymizedText(transformed, replacements)


class _FailingMappingStore(EncryptedMappingStore):
    def save(self, path, document, key):
        raise RuntimeError("encryption failed")


class _FailingValidator:
    def validate(self, source_path, output_path, mapping):
        raise RuntimeError("validation failed")


def _fail(source, destination):
    raise OSError("publish failed")


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()
