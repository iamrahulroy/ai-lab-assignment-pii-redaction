import hashlib
import re
import zipfile
from io import BytesIO
from pathlib import Path

from docx import Document
from lxml import etree
from PIL import Image

from pii_redactor.docx_adapter import DocxRedactor
from pii_redactor.domain import PseudonymizedText, TextReplacement


ORIGINAL = "Alice Example"
REPLACEMENT = "Morgan Reed"
REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"


def test_redacts_split_text_in_body_table_header_and_footer(tmp_path):
    source_path = tmp_path / "source.docx"
    output_path = tmp_path / "redacted.docx"
    document = Document()
    body = document.add_paragraph()
    body.add_run("Contact ").bold = True
    body.add_run("Alice ").italic = True
    body.add_run("Example").underline = True
    body.add_run(" today.")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).paragraphs[0].add_run("Alice ")
    table.cell(0, 0).paragraphs[0].add_run("Example")
    header = document.sections[0].header.paragraphs[0]
    header.add_run("Alice ")
    header.add_run("Example")
    footer = document.sections[0].footer.paragraphs[0]
    footer.add_run("Alice ")
    footer.add_run("Example")
    document.save(source_path)
    source_digest = hashlib.sha256(source_path.read_bytes()).hexdigest()

    DocxRedactor().redact(source_path, output_path, _pseudonymize)

    assert hashlib.sha256(source_path.read_bytes()).hexdigest() == source_digest
    redacted = Document(output_path)
    assert redacted.paragraphs[0].text == "Contact Morgan Reed today."
    assert redacted.tables[0].cell(0, 0).text == REPLACEMENT
    assert redacted.sections[0].header.paragraphs[0].text == REPLACEMENT
    assert redacted.sections[0].footer.paragraphs[0].text == REPLACEMENT
    assert len(redacted.paragraphs[0].runs) == 4
    assert redacted.paragraphs[0].runs[1].italic is True
    assert redacted.paragraphs[0].runs[2].underline is True


def test_removes_originals_from_hidden_parts_metadata_and_relationships(
    tmp_path,
):
    source_path = tmp_path / "hidden.docx"
    output_path = tmp_path / "sanitized.docx"
    document = Document()
    document.add_paragraph(f"Visible {ORIGINAL}")
    document.core_properties.author = ORIGINAL
    document.save(source_path)
    _inject_hidden_pii(source_path)

    DocxRedactor().redact(source_path, output_path, _pseudonymize)

    with zipfile.ZipFile(output_path) as package:
        assert "docProps/custom.xml" not in package.namelist()
        searchable = b"\n".join(
            package.read(name)
            for name in package.namelist()
            if name.endswith((".xml", ".rels"))
        )
        assert ORIGINAL.encode() not in searchable
        assert b"alice@example.com" not in searchable
        assert searchable.count(REPLACEMENT.encode()) >= 5
        assert b"custom-properties" not in package.read("_rels/.rels")
        assert b"/docProps/custom.xml" not in package.read(
            "[Content_Types].xml"
        )
        _assert_private_attributes_scrubbed(package)
        _assert_external_targets_neutralized(package)

    Document(output_path)


def test_replaces_every_supplied_image_without_changing_dimensions_or_links(
    tmp_path,
):
    source_path = Path(__file__).parents[1] / "Red Herring Prospectus (3).docx"
    output_path = tmp_path / "sanitized-prospectus.docx"
    source_digest = hashlib.sha256(source_path.read_bytes()).hexdigest()

    DocxRedactor().redact(source_path, output_path, _identity)

    assert hashlib.sha256(source_path.read_bytes()).hexdigest() == source_digest
    with (
        zipfile.ZipFile(source_path) as source,
        zipfile.ZipFile(output_path) as output,
    ):
        source_media = {
            name: source.read(name)
            for name in source.namelist()
            if name.startswith("word/media/")
        }
        output_media = {
            name: output.read(name)
            for name in output.namelist()
            if name.startswith("word/media/")
        }
        assert len(source_media) == 8
        assert output_media.keys() == source_media.keys()
        assert _image_relationships(output) == _image_relationships(source)
        for name, original_bytes in source_media.items():
            assert original_bytes not in output_media.values()
            assert _image_size(output_media[name]) == _image_size(original_bytes)

    assert zipfile.is_zipfile(output_path)
    Document(output_path)


def _pseudonymize(text: str) -> PseudonymizedText:
    replacements = tuple(
        TextReplacement(match.start(), match.end(), REPLACEMENT)
        for match in re.finditer(re.escape(ORIGINAL), text)
    )
    transformed = text.replace(ORIGINAL, REPLACEMENT)
    return PseudonymizedText(transformed, replacements)


def _identity(text: str) -> PseudonymizedText:
    return PseudonymizedText(text, ())


def _inject_hidden_pii(path: Path) -> None:
    additions = {
        "word/comments.xml": _story_xml("comments", "Comment Alice Example"),
        "word/footnotes.xml": _story_xml("footnotes", "Footnote Alice Example"),
        "word/endnotes.xml": _story_xml("endnotes", "Endnote Alice Example"),
        "docProps/custom.xml": (
            b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            b'<Properties xmlns="http://schemas.openxmlformats.org/'
            b'officeDocument/2006/custom-properties"><property name="Owner">'
            b'<value>Alice Example</value></property></Properties>'
        ),
    }
    with zipfile.ZipFile(path) as source:
        entries = {name: source.read(name) for name in source.namelist()}

    document_xml = entries["word/document.xml"]
    tracked = (
        b'<w:p w:rsidR="DEADBEEF"><w:del w:author="Alice Example" '
        b'w:date="2025-01-01T00:00:00Z"><w:r><w:delText>'
        b'Deleted Alice Example</w:delText></w:r></w:del></w:p>'
    )
    entries["word/document.xml"] = document_xml.replace(
        b"<w:sectPr", tracked + b"<w:sectPr", 1
    )
    entries["word/_rels/document.xml.rels"] = _append_relationship(
        entries["word/_rels/document.xml.rels"],
        "rIdExternalPii",
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        "mailto:alice@example.com",
        external=True,
    )
    entries["_rels/.rels"] = _append_relationship(
        entries["_rels/.rels"],
        "rIdCustomPii",
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/"
        "custom-properties",
        "docProps/custom.xml",
    )
    entries["[Content_Types].xml"] = entries["[Content_Types].xml"].replace(
        b"</Types>",
        b'<Override PartName="/docProps/custom.xml" '
        b'ContentType="application/vnd.openxmlformats-officedocument.'
        b'custom-properties+xml"/></Types>',
    )
    entries.update(additions)

    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as output:
        for name, data in entries.items():
            output.writestr(name, data)


def _story_xml(root_name: str, text: str) -> bytes:
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<w:{root_name} xmlns:w="http://schemas.openxmlformats.org/'
        f'wordprocessingml/2006/main"><w:p><w:r><w:t>{text}</w:t>'
        f'</w:r></w:p></w:{root_name}>'
    ).encode()


def _append_relationship(
    xml: bytes,
    relationship_id: str,
    relationship_type: str,
    target: str,
    external: bool = False,
) -> bytes:
    root = etree.fromstring(xml)
    attributes = {
        "Id": relationship_id,
        "Type": relationship_type,
        "Target": target,
    }
    if external:
        attributes["TargetMode"] = "External"
    etree.SubElement(root, f"{{{REL_NS}}}Relationship", **attributes)
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8")


def _assert_private_attributes_scrubbed(package: zipfile.ZipFile) -> None:
    for name in package.namelist():
        if not name.endswith(".xml"):
            continue
        root = etree.fromstring(package.read(name))
        for element in root.iter():
            for attribute, value in element.attrib.items():
                local_name = etree.QName(attribute).localname
                assert not local_name.startswith("rsid")
                if local_name in {"author", "initials"}:
                    assert value == ""


def _assert_external_targets_neutralized(package: zipfile.ZipFile) -> None:
    for name in package.namelist():
        if not name.endswith(".rels"):
            continue
        root = etree.fromstring(package.read(name))
        for relationship in root:
            if relationship.get("TargetMode") == "External":
                assert relationship.get("Target") == "https://example.invalid/"


def _image_relationships(package: zipfile.ZipFile) -> set[tuple[str, str]]:
    root = etree.fromstring(package.read("word/_rels/document.xml.rels"))
    return {
        (item.get("Id"), item.get("Target"))
        for item in root
        if item.get("Type", "").endswith("/image")
    }


def _image_size(data: bytes) -> tuple[int, int]:
    with Image.open(BytesIO(data)) as image:
        return image.size
