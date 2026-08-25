import hashlib
import zipfile

from docx import Document

from pii_redactor.bootstrap import build_redact_document


def test_redacts_a_styled_email_into_a_separate_docx(tmp_path):
    source_path = tmp_path / "source.docx"
    output_path = tmp_path / "redacted.docx"

    document = Document()
    paragraph = document.add_paragraph()
    lead = paragraph.add_run("Contact ")
    lead.bold = True
    email_start = paragraph.add_run("rashi.")
    email_start.italic = True
    email_end = paragraph.add_run("patil@example.com")
    email_end.underline = True
    tail = paragraph.add_run(" for details.")
    tail.bold = True
    document.add_paragraph("Email rashi.patil@example.com again.")
    document.save(source_path)

    source_digest = hashlib.sha256(source_path.read_bytes()).hexdigest()

    result = build_redact_document(seed=7).execute(source_path, output_path)

    assert hashlib.sha256(source_path.read_bytes()).hexdigest() == source_digest
    assert output_path != source_path
    assert zipfile.is_zipfile(output_path)

    redacted_document = Document(output_path)
    redacted_paragraph = redacted_document.paragraphs[0]
    redacted_text = "\n".join(item.text for item in redacted_document.paragraphs)
    assert "rashi.patil@example.com" not in redacted_text
    assert redacted_paragraph.text.startswith("Contact ")
    assert redacted_paragraph.text.endswith(" for details.")
    assert len(redacted_paragraph.runs) == 4
    assert redacted_paragraph.runs[0].bold is True
    assert redacted_paragraph.runs[1].italic is True
    assert redacted_paragraph.runs[2].underline is True
    assert redacted_paragraph.runs[3].bold is True

    assert len(result.mappings) == 1
    mapping = result.mappings[0]
    assert mapping.entity_type == "EMAIL_ADDRESS"
    assert mapping.original == "rashi.patil@example.com"
    assert mapping.replacement.endswith("@example.invalid")
    assert redacted_text.count(mapping.replacement) == 2
